#!/usr/bin/env python3
"""Клиентский отчёт о подсчёте пассажиров по машине за день: xlsx + письмо (українською).

Один отчёт — все камеры машины: строка на камеру и общий итог (сумма камер,
как aiValue на сайте). Чистые функции: ResultReport -> bytes / html / text.
Без БД и сети — данные собирает mailer.ResultMailer. Оформление повторяет
отчёты PassengersFlowCounting.

Время — локальное naive (Excel не поддерживает datetime с часовым поясом).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from html import escape
from io import BytesIO
from typing import List, Optional

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Ключи — job.SOURCE_AUTO / job.SOURCE_QUEUE.
SOURCE_LABELS = {
    "auto": "Автоматичний підрахунок за день",
    "queue": "Підрахунок за запитом",
}
# Что показывает число — как revenue3_resolveAutoCountInOut в revenue3.php.
MODE_DESCRIPTIONS = {
    "in": "кількість пасажирів, які увійшли до салону",
    "out": "кількість пасажирів, які вийшли із салону",
    "avr": "середнє між кількістю пасажирів, які увійшли до салону та вийшли з нього",
}


# ── DTO ──────────────────────────────────────────────────────────────

@dataclass
class CameraRow:
    cam_name: str
    window_start: Optional[datetime]
    window_end: Optional[datetime]
    passengers: int
    count_mode: str

    @property
    def period_label(self) -> str:
        return period_label(self.window_start, self.window_end)


@dataclass
class ResultReport:
    source: str
    car_name: str
    imei: str
    day: date
    version: int
    previous_sent_at: Optional[datetime]
    generated_at: datetime
    cameras: List[CameraRow] = field(default_factory=list)

    @property
    def passengers(self) -> int:
        return sum(c.passengers for c in self.cameras)

    @property
    def is_update(self) -> bool:
        return self.version > 1

    @property
    def source_label(self) -> str:
        return SOURCE_LABELS.get(self.source, self.source)

    @property
    def period_label(self) -> str:
        starts = [c.window_start for c in self.cameras if c.window_start]
        ends = [c.window_end for c in self.cameras if c.window_end]
        return period_label(min(starts) if starts else None, max(ends) if ends else None)

    @property
    def counted_description(self) -> str:
        modes = {c.count_mode for c in self.cameras}
        many = len(self.cameras) > 1
        if len(modes) == 1:
            what = MODE_DESCRIPTIONS.get(modes.pop(), MODE_DESCRIPTIONS["in"])
            text = (f"Підраховано {what}, за відеозаписами "
                    f"{'камер' if many else 'камери'} салону.")
        else:
            text = ("Для кожної камери салону підраховано кількість пасажирів, які "
                    "увійшли до салону, вийшли з нього або середнє між цими "
                    "значеннями — відповідно до налаштування камери.")
        if many:
            text += " Загальна кількість — сума по всіх камерах."
        return text

    @property
    def update_note(self) -> str:
        if not self.is_update:
            return ""
        note = (f"Це оновлена версія звіту (версія {self.version}): дані за цей день "
                "було перераховано. Вона замінює всі попередні версії")
        if self.previous_sent_at:
            note += f", остання з яких надіслана {fmt_datetime(self.previous_sent_at)}"
        return note + "."


# ── форматирование ───────────────────────────────────────────────────

def fmt_int(value) -> str:
    """1234 -> '1 234'."""
    return f"{int(value):,}".replace(",", "\u00a0")


def fmt_date(value) -> str:
    return value.strftime("%d.%m.%Y") if value else "—"


def fmt_datetime(value) -> str:
    return value.strftime("%d.%m.%Y %H:%M") if value else "—"


def period_label(start: Optional[datetime], end: Optional[datetime]) -> str:
    if not (start and end):
        return "—"
    end_fmt = "%H:%M" if start.date() == end.date() else "%d.%m.%Y %H:%M"
    return f"{fmt_datetime(start)} — {end.strftime(end_fmt)}"


def report_filename(report: ResultReport) -> str:
    """Звіт_2026-10-20_<авто>[_версія_2].xlsx"""
    car = re.sub(r"[^\w-]+", "_", report.car_name).strip("_")[:40] or report.imei
    name = f"Звіт_{report.day.isoformat()}_{car}"
    if report.is_update:
        name += f"_версія_{report.version}"
    return f"{name}.xlsx"


def email_subject(report: ResultReport) -> str:
    subject = (f"Звіт про підрахунок пасажирів за {fmt_date(report.day)} — "
               f"{report.car_name}")
    if report.is_update:
        subject += f" (версія {report.version}, оновлено)"
    return subject


# ── xlsx ─────────────────────────────────────────────────────────────

def build_xlsx(report: ResultReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    title_font = Font(size=14, bold=True)
    bold = Font(bold=True)
    big_bold = Font(size=12, bold=True)
    update_font = Font(bold=True, color="C62828")
    head_fill = PatternFill("solid", fgColor="E0F2F1")
    total_fill = PatternFill("solid", fgColor="F5F5F5")
    thin = Side(style="thin", color="BDBDBD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    wb.properties.title = f"Звіт про підрахунок пасажирів {fmt_date(report.day)}"
    wb.properties.creator = "Підрахунок пасажирів"

    ws = wb.active
    ws.title = "Звіт"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14

    def kv(row, label, value, font=None, number_format=None):
        ws.cell(row=row, column=1, value=label).font = bold
        cell = ws.cell(row=row, column=2, value=value)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if font is not None:
            cell.font = font
        if number_format:
            cell.number_format = number_format
        return row + 1

    def paragraph(row, label, text, font=None):
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=1).alignment = wrap
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        cell = ws.cell(row=row, column=2, value=text)
        cell.alignment = wrap
        if font is not None:
            cell.font = font
        ws.row_dimensions[row].height = 15 * max(1, -(-len(text) // 60)) + 3
        return row + 1

    title = "Звіт про підрахунок пасажирів"
    if report.is_update:
        title += f" (версія {report.version}, оновлено)"
    ws.cell(row=1, column=1, value=title).font = title_font

    row = 3
    row = kv(row, "Транспортний засіб", report.car_name)
    row = kv(row, "IMEI", report.imei)
    row = kv(row, "Дата підрахунку", fmt_date(report.day))
    row = kv(row, "Тип підрахунку", report.source_label)
    row = kv(row, "Період підрахунку", report.period_label)
    row = kv(row, "Версія звіту",
             f"{report.version} (оновлена)" if report.is_update else "1",
             font=update_font if report.is_update else None)
    if report.is_update:
        row = paragraph(row, "Увага", report.update_note, font=update_font)
    row = kv(row, "Звіт сформовано", fmt_datetime(report.generated_at))
    row += 1
    row = paragraph(row, "Що підраховано", report.counted_description)
    row += 1
    row = kv(row, "Пасажирів за день", report.passengers, font=big_bold,
             number_format="0")
    row += 1

    for idx, name in enumerate(("Камера", "Період підрахунку", "Пасажирів"), start=1):
        cell = ws.cell(row=row, column=idx, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = bold, head_fill, border, center
    row += 1
    for cam in report.cameras:
        for idx, value in enumerate((cam.cam_name, cam.period_label, cam.passengers),
                                    start=1):
            ws.cell(row=row, column=idx, value=value).border = border
        row += 1
    for idx, value in enumerate(("Разом", "", report.passengers), start=1):
        cell = ws.cell(row=row, column=idx, value=value)
        cell.font, cell.fill, cell.border = bold, total_fill, border

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── письмо ───────────────────────────────────────────────────────────

_TD_LABEL = "padding:8px 10px;border-bottom:1px solid #eeeeee;color:#616161;"
_TD_VALUE = "padding:8px 10px;border-bottom:1px solid #eeeeee;text-align:right;"
_TH = ("padding:8px 10px;border-bottom:2px solid #b2dfdb;color:#616161;"
       "font-weight:normal;text-align:left;")


def _html_row(label: str, value: str, extra_style: str = "") -> str:
    return (f'<tr><td style="{_TD_LABEL}">{escape(label)}</td>'
            f'<td style="{_TD_VALUE}{extra_style}">{value}</td></tr>')


def _html_cameras(report: ResultReport) -> str:
    rows = "".join(
        f'<tr><td style="{_TD_LABEL}color:#212121;">{escape(c.cam_name)}</td>'
        f'<td style="{_TD_LABEL}">{escape(c.period_label)}</td>'
        f'<td style="{_TD_VALUE}">{fmt_int(c.passengers)}</td></tr>'
        for c in report.cameras
    )
    return (
        '<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
        'style="border-collapse:collapse;margin:0 0 16px;font-size:13px;">'
        f'<tr><th style="{_TH}">Камера</th><th style="{_TH}">Період підрахунку</th>'
        f'<th style="{_TH}text-align:right;">Пасажирів</th></tr>{rows}</table>'
    )


def render_html(report: ResultReport, *, subject: str, filename: str,
                reply_to: str = "") -> str:
    e = escape
    update = ""
    if report.is_update:
        update = ('<p style="margin:0 0 16px;padding:12px 14px;background:#ffebee;'
                  'border-left:4px solid #c62828;color:#b71c1c;">'
                  f"{e(report.update_note)}</p>")
    rows = "".join([
        _html_row("Тип підрахунку", e(report.source_label)),
        _html_row("Період підрахунку", e(report.period_label)),
        _html_row("Пасажирів за день", f"<strong>{fmt_int(report.passengers)}</strong>",
                  "font-size:16px;"),
    ])
    footer = f"Лист сформовано автоматично {fmt_datetime(report.generated_at)}."
    if not reply_to:
        footer += " Будь ласка, не відповідайте на нього."
    version = f" — версія {report.version}" if report.is_update else ""
    return f"""<!DOCTYPE html>
<html lang="uk">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(subject)}</title>
</head>
<body style="margin:0;padding:0;background:#f4f6f8;font-family:Arial,Helvetica,sans-serif;color:#212121;">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f4f6f8;padding:24px 0;">
  <tr>
    <td align="center">
      <table role="presentation" width="600" cellpadding="0" cellspacing="0" style="max-width:600px;width:100%;background:#ffffff;border-radius:8px;overflow:hidden;">
        <tr>
          <td style="background:#009688;color:#ffffff;padding:20px 24px;font-size:18px;font-weight:bold;">
            Звіт про підрахунок пасажирів{version}
          </td>
        </tr>
        <tr>
          <td style="padding:24px;font-size:14px;line-height:1.5;">
            <p style="margin:0 0 12px;">Добрий день!</p>
            <p style="margin:0 0 16px;">
              Надсилаємо звіт про підрахунок пасажирів за <strong>{fmt_date(report.day)}</strong>
              по транспортному засобу <strong>{e(report.car_name)}</strong>.
            </p>
            {update}
            <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;margin:0 0 16px;">
              {rows}
            </table>
            {_html_cameras(report)}
            <p style="margin:0 0 12px;"><strong>Що підраховано.</strong> {e(report.counted_description)}</p>
            <p style="margin:16px 0 0;">
              Звіт також у вкладеному файлі <strong>{e(filename)}</strong>.
            </p>
          </td>
        </tr>
        <tr>
          <td style="padding:16px 24px;background:#fafafa;color:#757575;font-size:12px;line-height:1.4;">
            {e(footer)}
          </td>
        </tr>
      </table>
    </td>
  </tr>
</table>
</body>
</html>
"""


def render_text(report: ResultReport, *, filename: str, reply_to: str = "") -> str:
    lines = [
        "Добрий день!",
        "",
        f"Надсилаємо звіт про підрахунок пасажирів за {fmt_date(report.day)} "
        f"по транспортному засобу {report.car_name}.",
        "",
    ]
    if report.is_update:
        lines += [report.update_note, ""]
    lines += [
        f"Тип підрахунку: {report.source_label}",
        f"Період підрахунку: {report.period_label}",
        f"Пасажирів за день: {fmt_int(report.passengers)}",
        "",
        "По камерах:",
    ]
    lines += [f"  {c.cam_name} ({c.period_label}): {fmt_int(c.passengers)}"
              for c in report.cameras]
    lines += [
        "",
        f"Що підраховано. {report.counted_description}",
        "",
        f"Звіт також у вкладеному файлі {filename}.",
        "",
        "--",
        f"Лист сформовано автоматично {fmt_datetime(report.generated_at)}."
        + ("" if reply_to else " Будь ласка, не відповідайте на нього."),
    ]
    return "\n".join(lines) + "\n"
